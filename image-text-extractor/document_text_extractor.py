# ============================================================
# Document Text Extractor
# ============================================================
#
# Was macht dieses Programm?
#
# Dieses Programm extrahiert Text aus:
#
# - PDF-Dateien
# - Word-Dateien (.docx)
# - PowerPoint-Dateien (.pptx)
# - Bildern
#
# Unterstützte OCR-Sprachen:
#
# - Persisch / Farsi
# - Deutsch
# - Englisch
# - Persisch + Deutsch + Englisch gleichzeitig
#
# ------------------------------------------------------------
# PDF-Logik
# ------------------------------------------------------------
#
# 1. PDF enthält normalen Text:
#       -> Text wird direkt aus der PDF gelesen.
#
# 2. PDF ist ein Scan / besteht aus Bildern:
#       -> Ganze Seite wird als hochauflösendes Bild gerendert.
#       -> Bild wird verbessert.
#       -> Tesseract OCR erkennt den Text.
#
# 3. PDF enthält Text UND Bilder:
#       -> PDF-Text wird direkt übernommen.
#       -> Bilder auf derselben Seite werden zusätzlich mit
#          Tesseract OCR analysiert.
#
# ------------------------------------------------------------
# Word
# ------------------------------------------------------------
#
# - Normaler Word-Text wird direkt ausgelesen.
# - Tabellen werden ebenfalls gelesen.
# - Eingebettete Bilder werden zusätzlich per OCR untersucht.
#
# ------------------------------------------------------------
# PowerPoint
# ------------------------------------------------------------
#
# - Textfelder werden direkt gelesen.
# - Bilder auf den Folien werden zusätzlich per OCR untersucht.
#
# ------------------------------------------------------------
# Export
# ------------------------------------------------------------
#
# Das Ergebnis kann gespeichert werden als:
#
# - TXT
# - Word (.docx)
# - PowerPoint (.pptx)
# - mehrere Formate gleichzeitig
#
# ============================================================
# BENÖTIGTE PYTHON-BIBLIOTHEKEN
# ============================================================
#
# Einmal installieren:
#
# py -m pip install pytesseract pillow opencv-python python-docx python-pptx pymupdf
#
# ============================================================
# TESSERACT
# ============================================================
#
# Zusätzlich muss Tesseract OCR installiert sein.
#
# Benötigte Sprachen:
#
# eng = Englisch
# deu = Deutsch
# fas = Persisch
#
# Prüfen mit:
#
# tesseract --list-langs
#
# Falls Tesseract nicht im Windows PATH liegt, versucht
# dieses Programm automatisch typische Installationsorte.
#
# ============================================================


# ============================================================
# Standard-Bibliotheken
# ============================================================

import os
import shutil
import threading
import zipfile
import re
import subprocess
import tempfile

from io import BytesIO
from pathlib import Path


# ============================================================
# Tkinter
# ============================================================

import tkinter as tk

from tkinter import (
    ttk,
    filedialog,
    messagebox,
    simpledialog
)


# ============================================================
# Bilder
# ============================================================

from PIL import (
    Image,
    ImageTk,
    ImageOps
)


# ============================================================
# OpenCV / Bildverbesserung
# ============================================================

import cv2
import numpy as np


# ============================================================
# OCR
# ============================================================

import pytesseract


# ============================================================
# PDF
# ============================================================

import pymupdf


# ============================================================
# Word
# ============================================================

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.shared import Pt

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT


# ============================================================
# PowerPoint
# ============================================================

from pptx import Presentation

from pptx.util import Inches, Pt as PPTPt

from pptx.enum.text import PP_ALIGN

from pptx.enum.shapes import MSO_SHAPE_TYPE


# ============================================================
# Globale Variablen
# ============================================================

selected_files = []

extraction_results = []

# Neu: Seitenweise PDF-Texte und OCR-Hinweise für den bearbeitbaren
# Word-Export. Die bisherigen Textergebnisse bleiben unverändert.
pdf_page_results = {}

extraction_running = False

current_preview_image = None

# ============================================================
# Projektordner und eigene Tesseract-Sprachen
# ============================================================
#
# Dadurch verwenden Zuhause, Büro und andere Rechner
# immer dieselben Sprachdateien aus diesem Projekt.
# ============================================================

BASE_FOLDER = Path(__file__).resolve().parent

TESSDATA_FOLDER = BASE_FOLDER / "tessdata"


def tesseract_config(psm):

    # pytesseract zerlegt den Konfigurationsstring unter Windows.
    # Vorwärtsschrägstriche verhindern, dass C:\\... als Escape
    # interpretiert wird; Anführungszeichen erhalten Leerzeichen.
    data_dir = TESSDATA_FOLDER.resolve().as_posix()

    return (
        f'--tessdata-dir "{data_dir}" '
        f'--oem 3 --psm {psm}'
    )

# ============================================================
# Unterstützte Dateiendungen
# ============================================================


IMAGE_EXTENSIONS = {

    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp"

}


SUPPORTED_EXTENSIONS = (

    IMAGE_EXTENSIONS

    | {
        ".pdf",
        ".docx",
        ".pptx"
    }

)


# ============================================================
# Sprachen
# ============================================================

LANGUAGES = {

    "Persisch / Farsi":
        "fas",

    "Deutsch":
        "deu",

    "Englisch":
        "eng",

    "Persisch + Deutsch + Englisch":
        "fas+deu+eng"

}


# ============================================================
# Tesseract suchen
# ============================================================

def find_tesseract():

    # ========================================================
    # 1. Zuerst Windows PATH prüfen
    # ========================================================

    found = shutil.which(
        "tesseract"
    )

    if found:

        pytesseract.pytesseract.tesseract_cmd = found

        return True

    # ========================================================
    # 2. Verschiedene mögliche Installationen suchen
    # ========================================================

    possible_paths = [

        # ----------------------------------------------------
        # Installation wie auf deinem Büro-PC
        # ----------------------------------------------------

        Path.home()
        / "AppData"
        / "Local"
        / "Tesseract-OCR"
        / "tesseract.exe",


        # ----------------------------------------------------
        # Normale Windows-Installation
        # ----------------------------------------------------

        Path(
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        ),


        Path(
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"
        ),


        # ----------------------------------------------------
        # PDF24 Installation wie auf deinem Heim-PC
        # ----------------------------------------------------

        Path(
            r"C:\Program Files\PDF24\tesseract\tesseract.exe"
        )

    ]

    for path in possible_paths:

        if path.exists():

            pytesseract.pytesseract.tesseract_cmd = str(
                path
            )

            return True

    return False

# ============================================================
# Tesseract überprüfen
# ============================================================


def check_tesseract():

    # ========================================================
    # Tesseract.exe finden
    # ========================================================

    if not find_tesseract():

        messagebox.showerror(

            "Tesseract nicht gefunden",

            "Tesseract OCR wurde auf diesem Computer "
            "nicht gefunden.\n\n"
            "Bitte Tesseract installieren und danach "
            "das Programm erneut starten."

        )

        return False

    # ========================================================
    # Projektinternen tessdata-Ordner prüfen
    # ========================================================

    if not TESSDATA_FOLDER.exists():

        messagebox.showerror(

            "tessdata fehlt",

            "Der Sprachordner wurde nicht gefunden:\n\n"
            f"{TESSDATA_FOLDER}\n\n"
            "Der Ordner 'tessdata' muss neben "
            "'document_text_extractor.py' liegen."

        )

        return False

    # ========================================================
    # Gewählte Sprache bestimmen
    # ========================================================

    language_code = LANGUAGES[
        language_box.get()
    ]

    required_languages = language_code.split(
        "+"
    )

    # ========================================================
    # Prüfen, ob alle Sprachdateien vorhanden sind
    # ========================================================

    missing_languages = []

    for language in required_languages:

        language_file = (

            TESSDATA_FOLDER

            / f"{language}.traineddata"

        )

        if not language_file.exists():

            missing_languages.append(
                language
            )

    if missing_languages:

        messagebox.showerror(

            "Sprachdateien fehlen",

            "Folgende OCR-Sprachen fehlen im "
            "Projektordner:\n\n"
            + "\n".join(missing_languages)
            + "\n\n"
            f"Erwarteter Ordner:\n{TESSDATA_FOLDER}"

        )

        return False

    # Eine vorhandene Datei allein beweist nicht, dass Tesseract
    # sie am konfigurierten Pfad auch laden kann.
    try:

        pytesseract.image_to_string(
            Image.new("RGB", (160, 60), "white"),
            lang=language_code,
            config=tesseract_config(3),
            timeout=30
        )

    except Exception as error:

        messagebox.showerror(
            "OCR-Sprachtest fehlgeschlagen",
            "Tesseract konnte die gewählten Sprachdateien "
            "nicht laden.\n\n"
            f"Sprachordner: {TESSDATA_FOLDER.resolve()}\n"
            "Erwartet: "
            + ", ".join(
                f"{language}.traineddata"
                for language in required_languages
            )
            + "\n\n"
            "Bitte das komplette ZIP-Paket entpacken und "
            "prüfen, ob der Ordner 'tessdata' direkt neben "
            "dieser Python-Datei liegt.\n\n"
            f"Technischer Fehler:\n{error}"
        )

        return False

    return True

# ============================================================
# Dateien auswählen
# ============================================================


def add_files():

    files = filedialog.askopenfilenames(

        title="Dokumente auswählen",

        filetypes=[

            (
                "Unterstützte Dokumente",
                "*.pdf *.docx *.pptx "
                "*.jpg *.jpeg *.png *.bmp *.tif *.tiff *.webp"
            ),

            (
                "PDF",
                "*.pdf"
            ),

            (
                "Word",
                "*.docx"
            ),

            (
                "PowerPoint",
                "*.pptx"
            ),

            (
                "Bilder",
                "*.jpg *.jpeg *.png *.bmp *.tif *.tiff *.webp"
            ),

            (
                "Alle Dateien",
                "*.*"
            )

        ]

    )

    if not files:

        return

    unsupported = []

    for file_path in files:

        extension = Path(
            file_path
        ).suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:

            unsupported.append(
                Path(file_path).name
            )

            continue

        if file_path not in selected_files:

            selected_files.append(
                file_path
            )

    refresh_file_list()

    if unsupported:

        messagebox.showwarning(

            "Nicht unterstützt",

            "Folgende Dateien wurden ignoriert:\n\n"
            + "\n".join(unsupported)

        )

    status_label.config(

        text=f"{len(selected_files)} Datei(en) ausgewählt."

    )


# ============================================================
# Dateiliste aktualisieren
# ============================================================

def refresh_file_list():

    file_listbox.delete(
        0,
        tk.END
    )

    for index, file_path in enumerate(

        selected_files,

        start=1

    ):

        file_listbox.insert(

            tk.END,

            f"{index}. {Path(file_path).name}"

        )

    file_count_label.config(

        text=f"{len(selected_files)} Datei(en)"

    )


# ============================================================
# Datei entfernen
# ============================================================

def remove_file():

    selection = file_listbox.curselection()

    if not selection:

        return

    index = selection[0]

    selected_files.pop(
        index
    )

    refresh_file_list()

    preview_label.config(

        image="",

        text="Keine Vorschau"

    )

    preview_label.image = None


# ============================================================
# Alle Dateien entfernen
# ============================================================

def clear_files():

    if not selected_files:

        return

    answer = messagebox.askyesno(

        "Liste leeren",

        "Alle ausgewählten Dateien entfernen?"

    )

    if not answer:

        return

    selected_files.clear()

    refresh_file_list()

    preview_label.config(

        image="",

        text="Keine Vorschau"

    )

    preview_label.image = None

    status_label.config(
        text="Dateiliste wurde geleert."
    )


# ============================================================
# Datei nach oben
# ============================================================

def move_up():

    selection = file_listbox.curselection()

    if not selection:

        return

    index = selection[0]

    if index == 0:

        return

    selected_files[index - 1], selected_files[index] = (

        selected_files[index],

        selected_files[index - 1]

    )

    refresh_file_list()

    file_listbox.selection_set(
        index - 1
    )


# ============================================================
# Datei nach unten
# ============================================================

def move_down():

    selection = file_listbox.curselection()

    if not selection:

        return

    index = selection[0]

    if index >= len(selected_files) - 1:

        return

    selected_files[index + 1], selected_files[index] = (

        selected_files[index],

        selected_files[index + 1]

    )

    refresh_file_list()

    file_listbox.selection_set(
        index + 1
    )


# ============================================================
# Dateivorschau
# ============================================================

def file_selected(event=None):

    global current_preview_image

    selection = file_listbox.curselection()

    if not selection:

        return

    file_path = Path(

        selected_files[
            selection[0]
        ]

    )

    extension = file_path.suffix.lower()

    # ========================================================
    # Bild
    # ========================================================

    if extension in IMAGE_EXTENSIONS:

        try:

            image = Image.open(
                file_path
            )

            image = ImageOps.exif_transpose(
                image
            )

            show_preview_image(
                image
            )

        except Exception as error:

            preview_label.config(

                image="",

                text=f"Vorschau nicht möglich:\n{error}"

            )

    # ========================================================
    # PDF erste Seite
    # ========================================================

    elif extension == ".pdf":

        try:

            document = pymupdf.open(
                file_path
            )

            if len(document) == 0:

                document.close()

                return

            page = document[0]

            pixmap = page.get_pixmap(

                matrix=pymupdf.Matrix(
                    1.5,
                    1.5
                ),

                alpha=False

            )

            image = Image.open(

                BytesIO(
                    pixmap.tobytes("png")
                )

            )

            document.close()

            show_preview_image(
                image
            )

        except Exception as error:

            preview_label.config(

                image="",

                text=f"PDF-Vorschau nicht möglich:\n{error}"

            )

    # ========================================================
    # Word / PowerPoint
    # ========================================================

    else:

        preview_label.config(

            image="",

            text=(
                f"{file_path.name}\n\n"
                f"{extension.upper()[1:]} Dokument\n\n"
                "Text und eingebettete Bilder werden "
                "automatisch analysiert."
            )

        )

        preview_label.image = None


# ============================================================
# Vorschau anzeigen
# ============================================================

def show_preview_image(image):

    global current_preview_image

    preview = image.copy()

    preview.thumbnail(

        (
            390,
            300
        ),

        Image.Resampling.LANCZOS

    )

    photo = ImageTk.PhotoImage(
        preview
    )

    preview_label.config(

        image=photo,

        text=""

    )

    preview_label.image = photo

    current_preview_image = photo


# ============================================================
# PIL-Bild verbessern
# ============================================================

def preprocess_pil_image(
    pil_image,
    mode
):

    # EXIF-Drehung korrigieren
    pil_image = ImageOps.exif_transpose(
        pil_image
    )

    pil_image = pil_image.convert(
        "RGB"
    )

    # --------------------------------------------------------
    # Originalbild verwenden
    # --------------------------------------------------------

    if mode == "Original":

        return pil_image

    image = np.array(
        pil_image
    )

    # RGB -> Graustufen
    gray = cv2.cvtColor(

        image,

        cv2.COLOR_RGB2GRAY

    )

    height, width = gray.shape

    largest_side = max(
        width,
        height
    )

    # ========================================================
    # Kleine Bilder vergrößern
    # ========================================================

    if largest_side < 1200:

        scale = 2.5

    elif largest_side < 2000:

        scale = 2.0

    elif largest_side < 3000:

        scale = 1.5

    else:

        scale = 1.0

    if scale != 1.0:

        gray = cv2.resize(

            gray,

            None,

            fx=scale,

            fy=scale,

            interpolation=cv2.INTER_CUBIC

        )

    # ========================================================
    # Kontrast verbessern
    # ========================================================

    clahe = cv2.createCLAHE(

        clipLimit=2.0,

        tileGridSize=(8, 8)

    )

    gray = clahe.apply(
        gray
    )

    # ========================================================
    # Rauschen reduzieren
    # ========================================================

    gray = cv2.fastNlMeansDenoising(

        gray,

        None,

        h=7,

        templateWindowSize=7,

        searchWindowSize=21

    )

    # ========================================================
    # Automatisch
    # ========================================================

    if mode == "Automatisch":

        processed = cv2.adaptiveThreshold(

            gray,

            255,

            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,

            cv2.THRESH_BINARY,

            31,

            11

        )

    # ========================================================
    # Sanft
    # ========================================================

    elif mode == "Sanft":

        processed = gray

    # ========================================================
    # Stark
    # ========================================================

    else:

        _, processed = cv2.threshold(

            gray,

            0,

            255,

            cv2.THRESH_BINARY
            + cv2.THRESH_OTSU

        )

    # ========================================================
    # Weißer Rand
    # ========================================================

    processed = cv2.copyMakeBorder(

        processed,

        25,
        25,
        25,
        25,

        cv2.BORDER_CONSTANT,

        value=255

    )

    return Image.fromarray(
        processed
    )


# ============================================================
# OCR für PIL-Bild
# ============================================================

def ocr_pil_image(
    image,
    return_details=False
):

    language = LANGUAGES[
        language_box.get()
    ]

    processing_mode = (
        processing_box.get()
    )

    processed_image = preprocess_pil_image(

        image,

        processing_mode

    )

    # --------------------------------------------------------
    # PSM 3:
    # automatische Seitenerkennung
    # --------------------------------------------------------

    config = tesseract_config(3)

    text = pytesseract.image_to_string(

        processed_image,

        lang=language,

        config=config,

        timeout=180

    )

    text = clean_text(text)
    warnings = []

    # Bei zu wenig Text wird eine sanfte Bildversion mit PSM 6
    # versucht. Das Original wird nicht durch Fantasiezeichen ersetzt.
    if len(text) < 140:

        alternative = pytesseract.image_to_string(

            preprocess_pil_image(image, "Sanft"),

            lang=language,

            config=tesseract_config(6),

            timeout=180

        )

        alternative = clean_text(alternative)

        if len(alternative) > len(text):

            text = alternative
            warnings.append("OCR-Fallback mit PSM 6 verwendet")

    if len(text) < 40:

        warnings.append(
            "Sehr wenig Text erkannt; Seite am Scan prüfen"
        )

    if return_details:

        return text, warnings

    return text


# ============================================================
# Text bereinigen
# ============================================================

def clean_text(text):

    if not text:

        return ""

    lines = [

        line.rstrip()

        for line in text.splitlines()

    ]

    # Zu viele Leerzeilen vermeiden
    cleaned_lines = []

    previous_empty = False

    for line in lines:

        is_empty = (
            not line.strip()
        )

        if is_empty and previous_empty:

            continue

        cleaned_lines.append(
            line
        )

        previous_empty = is_empty

    return "\n".join(
        cleaned_lines
    ).strip()


# ============================================================
# Prüfen, ob Text ausreichend vorhanden ist
# ============================================================
#
# Das ist besonders wichtig bei PDFs.
#
# Eine Seite könnte nur:
#
# - Seitennummer
# - Wasserzeichen
# - Überschrift
#
# als echten PDF-Text enthalten, während der Hauptinhalt ein
# Scan ist.
#
# In diesem Fall soll trotzdem Vollseiten-OCR verwendet werden.
# ============================================================

def has_substantial_text(text):

    if not text:

        return False

    characters = [

        character

        for character in text

        if character.isalnum()

    ]

    return len(characters) >= 80


# ============================================================
# Bilddatei extrahieren
# ============================================================

def extract_from_image(
    file_path
):

    image = Image.open(
        file_path
    )

    text = ocr_pil_image(
        image
    )

    return (
        "===== OCR aus Bild =====\n\n"
        + (
            text
            if text
            else "[Kein Text erkannt]"
        )
    )


# ============================================================
# PDF-Seite in 300 DPI rendern
# ============================================================

def render_pdf_page(
    page
):

    dpi = 300

    zoom = (
        dpi / 72
    )

    pixmap = page.get_pixmap(

        matrix=pymupdf.Matrix(
            zoom,
            zoom
        ),

        alpha=False

    )

    image_bytes = pixmap.tobytes(
        "png"
    )

    image = Image.open(

        BytesIO(
            image_bytes
        )

    )

    image.load()

    return image


# ============================================================
# Eingebettete Bilder einer PDF-Seite OCR
# ============================================================

def extract_pdf_images_text(
    document,
    page
):

    results = []

    images = page.get_images(
        full=True
    )

    image_number = 0

    seen_xrefs = set()

    for image_info in images:

        xref = image_info[0]

        if xref in seen_xrefs:

            continue

        seen_xrefs.add(xref)

        # Eine OCR-Textebene liegt oft über einem ganzen Scanbild.
        # Dessen erneute OCR würde den Seiteninhalt verdoppeln.
        try:

            if any(
                rect.get_area() > page.rect.get_area() * 0.65
                for rect in page.get_image_rects(xref)
            ):

                continue

        except Exception:

            pass

        try:

            extracted = document.extract_image(
                xref
            )

            image_bytes = extracted.get(
                "image"
            )

            if not image_bytes:

                continue

            image = Image.open(

                BytesIO(
                    image_bytes
                )

            )

            # ------------------------------------------------
            # Sehr kleine Bilder sind oft:
            #
            # - Logos
            # - Symbole
            # - Linien
            #
            # Diese überspringen.
            # ------------------------------------------------

            width, height = image.size

            if (
                width < 120
                or height < 60
            ):

                continue

            image_number += 1

            text = ocr_pil_image(
                image
            )

            # Nur sinnvolle Ergebnisse übernehmen
            if len(text.strip()) >= 3:

                results.append(

                    (
                        image_number,
                        text
                    )

                )

        except Exception:

            # Ein kaputtes oder ungewöhnliches Bild
            # soll nicht die ganze PDF stoppen.
            continue

    return results


# ============================================================
# PDF extrahieren
# ============================================================

def extract_from_pdf(
    file_path
):

    document = pymupdf.open(
        file_path
    )

    complete_parts = []

    page_records = []

    total_pages = len(
        document
    )

    try:

        for page_index in range(
            total_pages
        ):

            page = document[
                page_index
            ]

            page_number = (
                page_index + 1
            )

            update_detail_status(

                f"PDF-Seite {page_number} "
                f"von {total_pages}"

            )

            # =================================================
            # Normalen PDF-Text lesen
            # =================================================

            native_text = clean_text(

                page.get_text(
                    "text"
                )

            )

            page_parts = [

                f"===== PDF Seite {page_number} ====="

            ]

            # =================================================
            # FALL A / C:
            #
            # Seite enthält ausreichend echten Text.
            #
            # Text direkt übernehmen und eingebettete Bilder
            # zusätzlich OCRen.
            # =================================================

            if has_substantial_text(
                native_text
            ):

                page_body = native_text

                page_warnings = []

                page_parts.append(
                    "\n[PDF-Text]\n"
                )

                page_parts.append(
                    native_text
                )

                image_results = (
                    extract_pdf_images_text(

                        document,

                        page

                    )
                )

                if image_results:

                    page_parts.append(
                        "\n[Text aus Bildern]\n"
                    )

                    for (
                        image_number,
                        image_text
                    ) in image_results:

                        page_parts.append(

                            f"\n--- Bild {image_number} ---\n"

                        )

                        page_parts.append(
                            image_text
                        )

                        page_body += (
                            f"\n\n[Text aus Bild {image_number}]\n"
                            + image_text
                        )

            # =================================================
            # FALL B:
            #
            # Kein oder fast kein echter Text.
            #
            # Ganze Seite als 300-DPI-Bild OCRen.
            # =================================================

            else:

                page_parts.append(
                    "\n[OCR der gesamten Seite]\n"
                )

                page_image = render_pdf_page(
                    page
                )

                ocr_text, page_warnings = ocr_pil_image(
                    page_image,
                    return_details=True
                )

                page_body = ocr_text

                if ocr_text:

                    page_parts.append(
                        ocr_text
                    )

                else:

                    page_parts.append(
                        "[Kein Text erkannt]"
                    )

            complete_parts.append(

                "\n".join(
                    page_parts
                )

            )

            page_records.append({
                "number": page_number,
                "text": page_body,
                "origin": (
                    "PDF-Text" if has_substantial_text(native_text)
                    else "OCR"
                ),
                "warnings": page_warnings
            })

    finally:

        document.close()

    pdf_page_results[str(Path(file_path).resolve())] = page_records

    return "\n\n".join(
        complete_parts
    )


# ============================================================
# Word: Tabellen auslesen
# ============================================================

def extract_word_tables(
    document
):

    table_parts = []

    for table_index, table in enumerate(

        document.tables,

        start=1

    ):

        rows = []

        for row in table.rows:

            values = [

                cell.text.strip()

                for cell in row.cells

            ]

            rows.append(

                " | ".join(
                    values
                )

            )

        if rows:

            table_parts.append(

                f"--- Tabelle {table_index} ---\n"
                + "\n".join(rows)

            )

    return table_parts


# ============================================================
# Word: eingebettete Bilder OCR
# ============================================================

def extract_word_images(
    file_path
):

    image_results = []

    # DOCX ist intern eine ZIP-Datei.
    with zipfile.ZipFile(
        file_path,
        "r"
    ) as archive:

        media_files = [

            name

            for name in archive.namelist()

            if name.startswith(
                "word/media/"
            )

        ]

        for index, media_name in enumerate(

            media_files,

            start=1

        ):

            try:

                image_bytes = archive.read(
                    media_name
                )

                image = Image.open(

                    BytesIO(
                        image_bytes
                    )

                )

                width, height = image.size

                if (
                    width < 120
                    or height < 60
                ):

                    continue

                text = ocr_pil_image(
                    image
                )

                if len(text.strip()) >= 3:

                    image_results.append(

                        (
                            index,
                            Path(media_name).name,
                            text
                        )

                    )

            except Exception:

                continue

    return image_results


# ============================================================
# Word extrahieren
# ============================================================

def extract_from_word(
    file_path
):

    document = Document(
        file_path
    )

    parts = [

        "===== Word Dokument ====="

    ]

    # ========================================================
    # Normaler Text
    # ========================================================

    paragraphs = [

        paragraph.text.strip()

        for paragraph in document.paragraphs

        if paragraph.text.strip()

    ]

    if paragraphs:

        parts.append(
            "\n[Word-Text]\n"
        )

        parts.append(

            "\n".join(
                paragraphs
            )

        )

    # ========================================================
    # Tabellen
    # ========================================================

    tables = extract_word_tables(
        document
    )

    if tables:

        parts.append(
            "\n[Tabellen]\n"
        )

        parts.append(

            "\n\n".join(
                tables
            )

        )

    # ========================================================
    # Bilder
    # ========================================================

    image_results = extract_word_images(
        file_path
    )

    if image_results:

        parts.append(
            "\n[Text aus eingebetteten Bildern]\n"
        )

        for (
            image_number,
            image_name,
            image_text
        ) in image_results:

            parts.append(

                f"\n--- Bild {image_number}: "
                f"{image_name} ---\n"

            )

            parts.append(
                image_text
            )

    if (
        not paragraphs
        and not tables
        and not image_results
    ):

        parts.append(
            "\n[Kein Text gefunden]"
        )

    return "\n".join(
        parts
    )


# ============================================================
# PowerPoint extrahieren
# ============================================================

def extract_from_powerpoint(
    file_path
):

    presentation = Presentation(
        file_path
    )

    result_parts = []

    total_slides = len(
        presentation.slides
    )

    for slide_index, slide in enumerate(

        presentation.slides,

        start=1

    ):

        update_detail_status(

            f"PowerPoint-Folie {slide_index} "
            f"von {total_slides}"

        )

        slide_parts = [

            f"===== PowerPoint Folie {slide_index} ====="

        ]

        slide_texts = []

        image_texts = []

        # ====================================================
        # Shapes durchsuchen
        # ====================================================

        for shape in slide.shapes:

            # ------------------------------------------------
            # Normaler Text
            # ------------------------------------------------

            if hasattr(
                shape,
                "has_text_frame"
            ):

                if shape.has_text_frame:

                    text = clean_text(
                        shape.text
                    )

                    if text:

                        slide_texts.append(
                            text
                        )

            # ------------------------------------------------
            # Bild
            # ------------------------------------------------

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:

                try:

                    image_bytes = shape.image.blob

                    image = Image.open(

                        BytesIO(
                            image_bytes
                        )

                    )

                    width, height = image.size

                    if (
                        width < 120
                        or height < 60
                    ):

                        continue

                    text = ocr_pil_image(
                        image
                    )

                    if len(text.strip()) >= 3:

                        image_texts.append(
                            text
                        )

                except Exception:

                    continue

        # ====================================================
        # Direkter Text
        # ====================================================

        if slide_texts:

            slide_parts.append(
                "\n[PowerPoint-Text]\n"
            )

            slide_parts.append(

                "\n\n".join(
                    slide_texts
                )

            )

        # ====================================================
        # OCR-Bilder
        # ====================================================

        if image_texts:

            slide_parts.append(
                "\n[Text aus Bildern]\n"
            )

            for image_index, image_text in enumerate(

                image_texts,

                start=1

            ):

                slide_parts.append(

                    f"\n--- Bild {image_index} ---\n"

                )

                slide_parts.append(
                    image_text
                )

        if (
            not slide_texts
            and not image_texts
        ):

            slide_parts.append(
                "\n[Kein Text gefunden]"
            )

        result_parts.append(

            "\n".join(
                slide_parts
            )

        )

    return "\n\n".join(
        result_parts
    )


# ============================================================
# Extraktion starten
# ============================================================

def start_extraction():

    global extraction_running

    if extraction_running:

        return

    if not selected_files:

        messagebox.showwarning(

            "Keine Dokumente",

            "Bitte zuerst mindestens ein Dokument auswählen."

        )

        return

    if not check_tesseract():

        return

    extraction_running = True

    extract_button.config(
        state="disabled"
    )

    save_button.config(
        state="disabled"
    )

    progress_bar["value"] = 0

    progress_bar["maximum"] = len(
        selected_files
    )

    result_text.delete(

        "1.0",

        tk.END

    )

    status_label.config(

        text="Textextraktion wird gestartet..."

    )

    threading.Thread(

        target=run_extraction,

        daemon=True

    ).start()


# ============================================================
# Alle Dateien verarbeiten
# ============================================================

def run_extraction():

    global extraction_results
    global extraction_running

    try:

        extraction_results = []

        pdf_page_results.clear()

        for index, file_path_string in enumerate(
            selected_files
        ):

            file_path = Path(
                file_path_string
            )

            extension = file_path.suffix.lower()

            window.after(

                0,

                status_label.config,

                {
                    "text":
                        f"Verarbeite Datei "
                        f"{index + 1} von "
                        f"{len(selected_files)}:\n"
                        f"{file_path.name}"
                }

            )

            # =================================================
            # Bild
            # =================================================

            if extension in IMAGE_EXTENSIONS:

                extracted_text = extract_from_image(
                    file_path
                )

            # =================================================
            # PDF
            # =================================================

            elif extension == ".pdf":

                extracted_text = extract_from_pdf(
                    file_path
                )

            # =================================================
            # Word
            # =================================================

            elif extension == ".docx":

                extracted_text = extract_from_word(
                    file_path
                )

            # =================================================
            # PowerPoint
            # =================================================

            elif extension == ".pptx":

                extracted_text = (
                    extract_from_powerpoint(
                        file_path
                    )
                )

            else:

                extracted_text = (
                    "[Dateiformat nicht unterstützt]"
                )

            extraction_results.append(

                {

                    "file":
                        str(file_path),

                    "text":
                        extracted_text

                }

            )

            window.after(

                0,

                progress_bar.config,

                {
                    "value":
                        index + 1
                }

            )

        window.after(

            0,

            extraction_finished

        )

    except Exception as error:

        window.after(

            0,

            extraction_failed,

            str(error)

        )

    finally:

        extraction_running = False


# ============================================================
# Detailstatus
# ============================================================

def update_detail_status(
    message
):

    window.after(

        0,

        detail_status_label.config,

        {
            "text":
                message
        }

    )


# ============================================================
# Fertig
# ============================================================

def extraction_finished():

    extract_button.config(
        state="normal"
    )

    save_button.config(
        state="normal"
    )

    detail_status_label.config(
        text=""
    )

    result_text.delete(

        "1.0",

        tk.END

    )

    for index, result in enumerate(

        extraction_results,

        start=1

    ):

        file_name = Path(
            result["file"]
        ).name

        result_text.insert(

            tk.END,

            "\n"
            + "=" * 70
            + "\n"

        )

        result_text.insert(

            tk.END,

            f"DATEI {index}: {file_name}\n"

        )

        result_text.insert(

            tk.END,

            "=" * 70
            + "\n\n"

        )

        result_text.insert(

            tk.END,

            result["text"]

        )

        result_text.insert(

            tk.END,

            "\n\n"

        )

    status_label.config(

        text=(
            f"Fertig. "
            f"{len(extraction_results)} Datei(en) verarbeitet."
        )

    )

    answer = messagebox.askyesno(

        "Extraktion abgeschlossen",

        "Die Textextraktion ist abgeschlossen.\n\n"
        "Möchtest du das Ergebnis jetzt speichern?"

    )

    if answer:

        show_export_dialog()


# ============================================================
# Fehler
# ============================================================

def extraction_failed(
    error
):

    extract_button.config(
        state="normal"
    )

    detail_status_label.config(
        text=""
    )

    status_label.config(
        text="Extraktion fehlgeschlagen."
    )

    messagebox.showerror(

        "Fehler",

        "Die Textextraktion konnte nicht "
        "abgeschlossen werden:\n\n"
        + error

    )


# ============================================================
# Bearbeiteten Text holen
# ============================================================

def get_edited_text():

    return result_text.get(

        "1.0",

        tk.END

    ).strip()


# ============================================================
# Text kopieren
# ============================================================

def copy_text():

    text = get_edited_text()

    if not text:

        messagebox.showwarning(

            "Kein Text",

            "Es ist kein Text vorhanden."

        )

        return

    window.clipboard_clear()

    window.clipboard_append(
        text
    )

    status_label.config(

        text="Text wurde in die Zwischenablage kopiert."

    )


# ============================================================
# Ergebnis löschen
# ============================================================

def clear_result():

    if not get_edited_text():

        return

    answer = messagebox.askyesno(

        "Text löschen",

        "Den gesamten erkannten Text wirklich löschen?"

    )

    if not answer:

        return

    result_text.delete(

        "1.0",

        tk.END

    )

    save_button.config(
        state="disabled"
    )


# ============================================================
# Prüfen, ob Persisch gewählt wurde
# ============================================================

def is_persian_mode():

    return "fas" in LANGUAGES[
        language_box.get()
    ]


# ============================================================
# Exportdialog
# ============================================================

def edited_pdf_pages():

    # Die vorhandenen Seitenmarkierungen bleiben beim Bearbeiten
    # im rechten Textfeld stehen. So gelangen Korrekturen auch in
    # den seitengetreuen Word-Export.
    text = get_edited_text()

    chunks = re.split(
        r"(?m)^===== PDF Seite (\d+) =====\s*$",
        text
    )

    expected = sum(
        len(pdf_page_results.get(str(Path(item["file"]).resolve()), []))
        for item in extraction_results
        if Path(item["file"]).suffix.lower() == ".pdf"
    )

    if (len(chunks) - 1) // 2 != expected:

        return None

    pages = []

    for index in range(1, len(chunks), 2):

        body = chunks[index + 1]

        # Der nächste Datei-Header ist kein Teil der PDF-Seite.
        body = re.split(r"(?m)^={70}\s*$", body, maxsplit=1)[0]

        body = re.sub(
            r"(?m)^\[(?:PDF-Text|OCR der gesamten Seite)\]\s*$",
            "",
            body
        )

        pages.append((int(chunks[index]), body.strip()))

    original_numbers = [
        page["number"]
        for item in extraction_results
        if Path(item["file"]).suffix.lower() == ".pdf"
        for page in pdf_page_results.get(
            str(Path(item["file"]).resolve()), []
        )
    ]

    if [number for number, _ in pages] != original_numbers:

        return None

    return [body for _, body in pages]


def add_word_text(paragraph, text, size, bold=False):

    set_word_rtl(paragraph)

    # Lateinische Wörter bleiben LTR innerhalb eines RTL-Absatzes.
    for part in re.split(r"([A-Za-z][A-Za-z0-9 ._-]*)", text):

        if not part:

            continue

        run = paragraph.add_run(part)
        run.font.name = "DejaVu Sans"
        run.font.size = Pt(size)
        run.bold = bold

        direction = OxmlElement("w:rtl")
        direction.set(
            qn("w:val"),
            "0" if re.match(r"[A-Za-z]", part) else "1"
        )
        run._element.get_or_add_rPr().append(direction)


def add_writing_line(paragraph):

    border = OxmlElement("w:pBdr")

    for name in ("bottom", "between"):

        edge = OxmlElement("w:" + name)

        for key, value in (
            ("val", "single"),
            ("sz", "4"),
            ("color", "999999")
        ):

            edge.set(qn("w:" + key), value)

        border.append(edge)

    paragraph._p.get_or_add_pPr().append(border)


def save_pdf_as_styled_word(source_path, pages, file_path):

    document = Document()
    section = document.sections[0]

    # Format und Notizspalte entsprechen der ersten Lektion.
    section.page_width = Pt(612)
    section.page_height = Pt(792)
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(32.4)
    section.left_margin = Pt(39.6)
    section.right_margin = Pt(39.6)
    section.footer_distance = Pt(24)

    for index, page in enumerate(pages):

        if index:

            document.add_page_break()

        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Pt(72)
        table.columns[1].width = Pt(464.4)

        notes, body = table.rows[0].cells
        notes.width = Pt(72)
        body.width = Pt(464.4)

        cell_borders = OxmlElement("w:tcBorders")
        edge = OxmlElement("w:right")

        for key, value in (
            ("val", "single"),
            ("sz", "6"),
            ("color", "888888")
        ):

            edge.set(qn("w:" + key), value)

        cell_borders.append(edge)
        notes._tc.get_or_add_tcPr().append(cell_borders)

        notes.text = ""
        add_word_text(notes.paragraphs[0], "یادداشت ها", 15, True)
        notes.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        body.text = ""
        add_word_text(
            body.paragraphs[0],
            source_path.stem,
            20,
            True
        )
        body.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        blocks = [
            block.strip()
            for block in re.split(r"\n\s*\n", page["text"])
            if block.strip()
        ]

        density = len(page["text"].splitlines())
        font_size = 14.5 if density > 58 else 16 if density > 45 else 18.5

        for block_index, block in enumerate(blocks):

            paragraph = body.add_paragraph()

            heading = (
                len(block) < 100
                and (
                    block_index < 2
                    or block.startswith(
                        ("بخش ", "مطالعه ", "الف-", "ب-")
                    )
                )
            )

            add_word_text(
                paragraph,
                re.sub(r"\s*\n\s*", " ", block),
                font_size + 0.5 if heading else font_size,
                heading
            )

            paragraph.paragraph_format.space_after = Pt(
                3 if heading else 1
            )

        # Bei kurzen Übungsseiten einen editierbaren Schreibbereich.
        if len(blocks) < 10:

            for _ in range(8):

                paragraph = body.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(14)
                add_writing_line(paragraph)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    document.save(file_path)


def render_word_for_review(file_path, original_count):

    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    if not soffice:

        return "LibreOffice fehlt: Word-Seiten nicht gerendert."

    folder = file_path.parent / (file_path.stem + "_Pruefung")
    folder.mkdir(exist_ok=True)

    with tempfile.TemporaryDirectory() as temp_profile:

        profile_uri = Path(temp_profile).resolve().as_uri()

        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={profile_uri}",
                "--headless", "--convert-to", "pdf",
                "--outdir", str(folder), str(file_path)
            ],
            capture_output=True,
            text=True,
            timeout=180,
            check=True
        )

    preview_pdf = folder / (file_path.stem + ".pdf")

    if not preview_pdf.exists():

        raise RuntimeError("LibreOffice hat keine PDF-Vorschau erzeugt.")

    with pymupdf.open(preview_pdf) as rendered:

        count = len(rendered)

        for index, page in enumerate(rendered, 1):

            image = page.get_pixmap(
                matrix=pymupdf.Matrix(1.4, 1.4),
                alpha=False
            )

            image.save(folder / f"Seite-{index:03}.png")

    suffix = (
        " – Seitenzahl abweichend: Layout prüfen"
        if count != original_count else ""
    )

    return (
        f"{file_path.name}: {count} Word-Seiten, "
        f"{original_count} PDF-Seiten{suffix}; "
        f"Seitenbilder in {folder.name}"
    )


def export_pdf_review(output_folder, base_name, render_pages):

    pdf_items = [
        item for item in extraction_results
        if Path(item["file"]).suffix.lower() == ".pdf"
    ]

    if not pdf_items:

        raise ValueError("Für diesen Export muss eine PDF ausgewählt sein.")

    edited = edited_pdf_pages()
    notes = [
        "Prüfbericht: OCR ist keine bestätigte 1:1-Abschrift.",
        "Namen, Zahlen, Bibelstellen und schwer lesbare Wörter am Original prüfen.",
        ""
    ]

    if edited is None:

        notes.append(
            "Die PDF-Seitenmarkierungen im Textfeld wurden verändert. "
            "Der stilisierte Word-Export verwendet den ursprünglichen "
            "seitenweisen Text."
        )

    saved = []
    position = 0

    for file_index, item in enumerate(pdf_items, 1):

        source = Path(item["file"])
        originals = pdf_page_results[str(source.resolve())]
        pages = []

        for original in originals:

            page = dict(original)

            if edited is not None:

                page["text"] = edited[position]

            position += 1
            pages.append(page)

            for warning in page["warnings"]:

                notes.append(
                    f"{source.name}, Seite {page['number']}: {warning}"
                )

        suffix = f"_{file_index}" if len(pdf_items) > 1 else ""
        target = output_folder / f"{base_name}{suffix}_bearbeitbar.docx"

        save_pdf_as_styled_word(source, pages, target)
        saved.append(target.name)

        if render_pages:

            notes.append(
                render_word_for_review(target, len(pages))
            )

    report = output_folder / f"{base_name}_Pruefbericht.txt"
    report.write_text("\n".join(notes), encoding="utf-8-sig")
    saved.append(report.name)

    return saved


def show_export_dialog():

    if not get_edited_text():

        messagebox.showwarning(

            "Kein Text",

            "Es gibt keinen Text zum Speichern."

        )

        return

    dialog = tk.Toplevel(
        window
    )

    dialog.title(
        "Ergebnis speichern"
    )

    dialog.geometry(
        "480x510"
    )

    dialog.resizable(
        False,
        False
    )

    dialog.transient(
        window
    )

    dialog.grab_set()

    tk.Label(

        dialog,

        text="Exportformat auswählen",

        font=(
            "Segoe UI",
            16,
            "bold"
        )

    ).pack(
        pady=(25, 5)
    )

    tk.Label(

        dialog,

        text=(
            "Du kannst ein oder mehrere Formate "
            "gleichzeitig auswählen."
        )

    ).pack(
        pady=(0, 20)
    )

    txt_var = tk.BooleanVar(
        value=True
    )

    word_var = tk.BooleanVar(
        value=False
    )

    ppt_var = tk.BooleanVar(
        value=False
    )

    styled_word_var = tk.BooleanVar(value=False)

    render_word_var = tk.BooleanVar(value=False)

    tk.Checkbutton(

        dialog,

        text="Textdatei (.txt)",

        variable=txt_var,

        font=(
            "Segoe UI",
            11
        )

    ).pack(

        anchor="w",

        padx=90,

        pady=8

    )

    tk.Checkbutton(

        dialog,

        text="Word-Datei (.docx)",

        variable=word_var,

        font=(
            "Segoe UI",
            11
        )

    ).pack(

        anchor="w",

        padx=90,

        pady=8

    )

    tk.Checkbutton(

        dialog,

        text="PowerPoint (.pptx)",

        variable=ppt_var,

        font=(
            "Segoe UI",
            11
        )

    ).pack(

        anchor="w",

        padx=90,

        pady=8

    )

    tk.Checkbutton(
        dialog,
        text="PDF → bearbeitbares Word im Stil der ersten Lektion",
        variable=styled_word_var,
        font=("Segoe UI", 10)
    ).pack(anchor="w", padx=35, pady=7)

    tk.Checkbutton(
        dialog,
        text="Word-Seiten als Bilder zur Prüfung rendern",
        variable=render_word_var,
        font=("Segoe UI", 10)
    ).pack(anchor="w", padx=35, pady=7)

    def save_selected():

        formats = []

        if txt_var.get():

            formats.append(
                "txt"
            )

        if word_var.get():

            formats.append(
                "docx"
            )

        if ppt_var.get():

            formats.append(
                "pptx"
            )

        if styled_word_var.get():

            formats.append("styled_docx")

        if render_word_var.get():

            formats.append("render_word")

        if "render_word" in formats and "styled_docx" not in formats:

            messagebox.showwarning(
                "Word-Seiten rendern",
                "Bitte auch den PDF-zu-Word-Export auswählen.",
                parent=dialog
            )

            return

        if not formats:

            messagebox.showwarning(

                "Kein Format",

                "Bitte mindestens ein Format auswählen.",

                parent=dialog

            )

            return

        dialog.destroy()

        export_results(
            formats
        )

    button_frame = tk.Frame(
        dialog
    )

    button_frame.pack(
        pady=30
    )

    tk.Button(

        button_frame,

        text="Speichern",

        command=save_selected,

        width=15,

        height=2

    ).grid(

        row=0,
        column=0,
        padx=8

    )

    tk.Button(

        button_frame,

        text="Abbrechen",

        command=dialog.destroy,

        width=15,

        height=2

    ).grid(

        row=0,
        column=1,
        padx=8

    )


# ============================================================
# Export durchführen
# ============================================================

def export_results(
    formats
):

    output_folder = filedialog.askdirectory(

        title="Speicherordner auswählen"

    )

    if not output_folder:

        return

    base_name = simpledialog.askstring(

        "Dateiname",

        "Wie sollen die Dateien heißen?",

        initialvalue="Dokument_Text"

    )

    if not base_name:

        return

    # Ungültige Windows-Zeichen ersetzen
    invalid_characters = (
        '<>:"/\\|?*'
    )

    for character in invalid_characters:

        base_name = base_name.replace(

            character,

            "_"

        )

    output_folder = Path(
        output_folder
    )

    saved_files = []

    try:

        # ====================================================
        # TXT
        # ====================================================

        if "txt" in formats:

            file_path = (

                output_folder

                / f"{base_name}.txt"

            )

            save_as_txt(
                file_path
            )

            saved_files.append(
                file_path.name
            )

        # ====================================================
        # Word
        # ====================================================

        if "docx" in formats:

            file_path = (

                output_folder

                / f"{base_name}.docx"

            )

            save_as_word(
                file_path
            )

            saved_files.append(
                file_path.name
            )

        # ====================================================
        # PowerPoint
        # ====================================================

        if "pptx" in formats:

            file_path = (

                output_folder

                / f"{base_name}.pptx"

            )

            save_as_powerpoint(
                file_path
            )

            saved_files.append(
                file_path.name
            )

        # Neu: Für jede PDF ein eigener bearbeitbarer Word-Export,
        # plus Prüfbericht und optional gerenderte Seitenbilder.
        if "styled_docx" in formats:

            saved_files.extend(
                export_pdf_review(
                    output_folder,
                    base_name,
                    "render_word" in formats
                )
            )

        answer = messagebox.askyesno(

            "Gespeichert",

            "Folgende Dateien wurden erstellt:\n\n"
            + "\n".join(saved_files)
            + "\n\n"
            "Speicherordner öffnen?"

        )

        if answer:

            if hasattr(os, "startfile"):

                os.startfile(output_folder)

    except Exception as error:

        messagebox.showerror(

            "Export Fehler",

            f"Die Dateien konnten nicht gespeichert werden:\n\n"
            f"{error}"

        )


# ============================================================
# TXT speichern
# ============================================================

def save_as_txt(
    file_path
):

    # utf-8-sig ist auf Windows praktisch,
    # besonders bei persischen Zeichen.

    with open(

        file_path,

        "w",

        encoding="utf-8-sig"

    ) as file:

        file.write(
            get_edited_text()
        )


# ============================================================
# Word RTL
# ============================================================

def set_word_rtl(
    paragraph
):

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    paragraph_properties = (
        paragraph._p.get_or_add_pPr()
    )

    bidi = OxmlElement(
        "w:bidi"
    )

    paragraph_properties.append(
        bidi
    )


# ============================================================
# Word speichern
# ============================================================

def save_as_word(
    file_path
):

    document = Document()

    title = document.add_heading(

        "Extrahierter Dokumenttext",

        level=0

    )

    if is_persian_mode():

        set_word_rtl(
            title
        )

    text = get_edited_text()

    for line in text.splitlines():

        paragraph = document.add_paragraph()

        if is_persian_mode():

            set_word_rtl(
                paragraph
            )

        run = paragraph.add_run(
            line
        )

        run.font.name = "Arial"

        run.font.size = Pt(
            12
        )

    document.save(
        file_path
    )


# ============================================================
# Text für PowerPoint aufteilen
# ============================================================

def split_text_for_slides(
    text,
    max_characters=1300
):

    lines = text.splitlines()

    chunks = []

    current = ""

    for line in lines:

        proposed = (

            current
            + "\n"
            + line

        ).strip()

        if (
            len(proposed)
            > max_characters
        ):

            if current:

                chunks.append(
                    current.strip()
                )

            current = line

        else:

            current = proposed

    if current:

        chunks.append(
            current.strip()
        )

    return chunks


# ============================================================
# PowerPoint speichern
# ============================================================

def save_as_powerpoint(
    file_path
):

    presentation = Presentation()

    presentation.slide_width = Inches(
        13.333
    )

    presentation.slide_height = Inches(
        7.5
    )

    # ========================================================
    # Titelfolie
    # ========================================================

    title_slide = presentation.slides.add_slide(

        presentation.slide_layouts[0]

    )

    title_slide.shapes.title.text = (

        "Extrahierter Dokumenttext"

    )

    title_slide.placeholders[1].text = (

        f"{len(selected_files)} "
        f"Quelldatei(en)"

    )

    # ========================================================
    # Textfolien
    # ========================================================

    chunks = split_text_for_slides(

        get_edited_text()

    )

    for index, chunk in enumerate(

        chunks,

        start=1

    ):

        slide = presentation.slides.add_slide(

            presentation.slide_layouts[5]

        )

        slide.shapes.title.text = (

            f"Text {index} von {len(chunks)}"

        )

        text_box = slide.shapes.add_textbox(

            Inches(0.7),

            Inches(1.4),

            Inches(11.9),

            Inches(5.5)

        )

        text_frame = text_box.text_frame

        text_frame.word_wrap = True

        paragraph = (
            text_frame.paragraphs[0]
        )

        paragraph.text = chunk

        if is_persian_mode():

            paragraph.alignment = (
                PP_ALIGN.RIGHT
            )

        else:

            paragraph.alignment = (
                PP_ALIGN.LEFT
            )

        for run in paragraph.runs:

            run.font.name = "Arial"

            run.font.size = PPTPt(
                18
            )

    presentation.save(
        file_path
    )


# ============================================================
# Hauptfenster
# ============================================================

window = tk.Tk()


window.title(
    "Document Text Extractor"
)


window.geometry(
    "1200x820"
)


window.minsize(
    1000,
    720
)


# ============================================================
# Titel
# ============================================================

title_label = tk.Label(

    window,

    text="Document Text Extractor",

    font=(
        "Segoe UI",
        20,
        "bold"
    )

)


title_label.pack(
    pady=(15, 3)
)


description_label = tk.Label(

    window,

    text=(
        "Text aus PDF, Word, PowerPoint "
        "und Bildern extrahieren"
    )

)


description_label.pack(
    pady=(0, 10)
)


# ============================================================
# Hauptbereich
# ============================================================

main_frame = tk.Frame(
    window
)


main_frame.pack(

    fill="both",

    expand=True,

    padx=15,

    pady=5

)


# ============================================================
# Linke Seite
# ============================================================

left_frame = tk.Frame(

    main_frame,

    width=440

)


left_frame.pack(

    side="left",

    fill="y",

    padx=(0, 10)

)


# ============================================================
# Dateien auswählen
# ============================================================

add_button = tk.Button(

    left_frame,

    text="Dokumente auswählen",

    command=add_files,

    width=26,

    height=2

)


add_button.pack(
    pady=5
)


file_count_label = tk.Label(

    left_frame,

    text="0 Dateien"

)


file_count_label.pack(
    pady=3
)


# ============================================================
# Dateiliste
# ============================================================

list_frame = tk.Frame(
    left_frame
)


list_frame.pack(
    pady=5
)


file_listbox = tk.Listbox(

    list_frame,

    width=48,

    height=9,

    font=(
        "Segoe UI",
        9
    )

)


file_listbox.pack(

    side="left",

    fill="y"

)


file_scrollbar = tk.Scrollbar(

    list_frame,

    command=file_listbox.yview

)


file_scrollbar.pack(

    side="right",

    fill="y"

)


file_listbox.config(

    yscrollcommand=file_scrollbar.set

)


file_listbox.bind(

    "<<ListboxSelect>>",

    file_selected

)


# ============================================================
# Listenbuttons
# ============================================================

file_buttons = tk.Frame(
    left_frame
)


file_buttons.pack(
    pady=5
)


tk.Button(

    file_buttons,

    text="↑",

    command=move_up,

    width=5

).grid(

    row=0,
    column=0,
    padx=3

)


tk.Button(

    file_buttons,

    text="↓",

    command=move_down,

    width=5

).grid(

    row=0,
    column=1,
    padx=3

)


tk.Button(

    file_buttons,

    text="Entfernen",

    command=remove_file,

    width=10

).grid(

    row=0,
    column=2,
    padx=3

)


tk.Button(

    file_buttons,

    text="Alles löschen",

    command=clear_files,

    width=12

).grid(

    row=0,
    column=3,
    padx=3

)


# ============================================================
# Vorschau
# ============================================================

preview_frame = tk.Frame(

    left_frame,

    width=410,

    height=300,

    relief="groove",

    borderwidth=2

)


preview_frame.pack(
    pady=10
)


preview_frame.pack_propagate(
    False
)


preview_label = tk.Label(

    preview_frame,

    text="Keine Vorschau"

)


preview_label.pack(

    fill="both",

    expand=True

)


# ============================================================
# Einstellungen
# ============================================================

settings_frame = tk.LabelFrame(

    left_frame,

    text="Texterkennung",

    padx=10,

    pady=10

)


settings_frame.pack(

    fill="x",

    pady=8

)


# ------------------------------------------------------------
# Sprache
# ------------------------------------------------------------

tk.Label(

    settings_frame,

    text="Sprache:"

).grid(

    row=0,
    column=0,

    sticky="w",

    padx=5,
    pady=5

)


language_box = ttk.Combobox(

    settings_frame,

    values=list(
        LANGUAGES.keys()
    ),

    state="readonly",

    width=29

)


language_box.grid(

    row=0,
    column=1,

    padx=5,
    pady=5

)


language_box.set(
    "Persisch / Farsi"
)


# ------------------------------------------------------------
# Bildverbesserung
# ------------------------------------------------------------

tk.Label(

    settings_frame,

    text="Bildverbesserung:"

).grid(

    row=1,
    column=0,

    sticky="w",

    padx=5,
    pady=5

)


processing_box = ttk.Combobox(

    settings_frame,

    values=[

        "Automatisch",
        "Sanft",
        "Stark",
        "Original"

    ],

    state="readonly",

    width=29

)


processing_box.grid(

    row=1,
    column=1,

    padx=5,
    pady=5

)


processing_box.set(
    "Automatisch"
)


# ============================================================
# Start
# ============================================================

extract_button = tk.Button(

    left_frame,

    text="Text extrahieren",

    command=start_extraction,

    width=26,

    height=2

)


extract_button.pack(
    pady=10
)


# ============================================================
# Fortschritt
# ============================================================

progress_bar = ttk.Progressbar(

    left_frame,

    length=380,

    mode="determinate"

)


progress_bar.pack(
    pady=5
)


detail_status_label = tk.Label(

    left_frame,

    text="",

    wraplength=390,

    font=(
        "Segoe UI",
        9
    )

)


detail_status_label.pack(
    pady=3
)


# ============================================================
# Rechte Seite
# ============================================================

right_frame = tk.Frame(
    main_frame
)


right_frame.pack(

    side="right",

    fill="both",

    expand=True

)


result_title = tk.Label(

    right_frame,

    text="Extrahierter Text",

    font=(
        "Segoe UI",
        12,
        "bold"
    )

)


result_title.pack(
    pady=(5, 5)
)


# ============================================================
# Textfeld
# ============================================================

text_frame = tk.Frame(
    right_frame
)


text_frame.pack(

    fill="both",

    expand=True

)


result_text = tk.Text(

    text_frame,

    wrap="word",

    font=(
        "Arial",
        12
    ),

    undo=True

)


result_text.pack(

    side="left",

    fill="both",

    expand=True

)


text_scrollbar = tk.Scrollbar(

    text_frame,

    command=result_text.yview

)


text_scrollbar.pack(

    side="right",

    fill="y"

)


result_text.config(

    yscrollcommand=text_scrollbar.set

)


# ============================================================
# Ergebnisbuttons
# ============================================================

result_buttons = tk.Frame(
    right_frame
)


result_buttons.pack(
    pady=10
)


save_button = tk.Button(

    result_buttons,

    text="Ergebnis speichern",

    command=show_export_dialog,

    width=19,

    height=2,

    state="disabled"

)


save_button.grid(

    row=0,
    column=0,
    padx=5

)


copy_button = tk.Button(

    result_buttons,

    text="Text kopieren",

    command=copy_text,

    width=15,

    height=2

)


copy_button.grid(

    row=0,
    column=1,
    padx=5

)


clear_result_button = tk.Button(

    result_buttons,

    text="Text löschen",

    command=clear_result,

    width=15,

    height=2

)


clear_result_button.grid(

    row=0,
    column=2,
    padx=5

)


# ============================================================
# Status
# ============================================================

status_label = tk.Label(

    window,

    text="Dokumente auswählen, um zu beginnen.",

    wraplength=1100

)


status_label.pack(
    pady=(5, 15)
)


# ============================================================
# Programm starten
# ============================================================

window.mainloop()
